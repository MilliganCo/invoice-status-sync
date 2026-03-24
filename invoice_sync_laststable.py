# -*- coding: utf-8 -*-
"""
invoice_sync_v15.py
Обновление статусов по инвойсам без сдвига «якорных» колонок.

Изменения:
- Умный выбор листа статуса: предпочитаем видимый «ПРОЕКТ», затем видимый «ДЛЯ МОДЖИ»,
  иначе — лист с лучшим совпадением по заголовкам («артикул», «размер», «отгруз(ка)»).
- Доп. отладочный вывод: какой лист выбран, границы зоны инвойсов.

Сохранено:
- Адаптивная зона инвойсов (включая сценарий «после СУММА/ИТОГ/РАЗНИЦА»).
- Нормализация размеров (включая 164-84-92).
- Формулы «ОБЩЕЕ/ИТОГ/ВСЕГО» по каждому артикулу.
- В режиме --watch: антидубликаты (in-memory + processed_invoices.txt), проверка стабильности файла,
  ретраи сохранения при PermissionError, ограничение сканирования по фактической границе данных.
"""

import argparse, re, time, os
from copy import copy
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Tuple, Optional

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell.cell import MergedCell
from threading import Timer

# --- watchdog (опционально) ---
try:
    from watchdog.observers import Observer
    from watchdog.events import FileSystemEventHandler
    WATCHDOG_AVAILABLE = True
except Exception:
    # Watchdog not available. Provide a minimal stub so class definition below doesn't fail.
    WATCHDOG_AVAILABLE = False
    class FileSystemEventHandler:  # type: ignore
        pass

CONFIG = {
    "INVOICES_ROOT": r"C:\example\invoices\automation",
    "STATUS_ROOT":  r"C:\example\statuses\automation",
    "OUTPUT_SUBFOLDER": "_готово",
    # Входящие инвойсы допускают также .xls (для авто-конвертации)
    "ALLOWED_INVOICE_EXT": (".xlsx", ".xlsm", ".xls"),
    # Файлы статусов поддерживаются openpyxl: только .xlsx/.xlsm
    "ALLOWED_STATUS_EXT": (".xlsx", ".xlsm"),
    "FILE_SETTLE_SECONDS": 2,
    "SCAN_MAX_ROWS": 120,
    "SCAN_MAX_COLS": 200,

    # Антидубликаты watcher:
    "WATCH_DEDUP_COOLDOWN": 300,
    "WATCH_STABLE_WAIT": 1.0,
    "SAVE_RETRIES": 5,
    "SAVE_RETRY_BASE_DELAY": 1.0,
    # Конвертация .xls -> .xlsxа
    "DELETE_XLS_AFTER_CONVERT": True,
    "DELETE_XLS_DELAY": 5.0,
}

def log(msg: str): print(msg)

# -------------------- utils --------------------

def norm(s: Optional[str]) -> str:
    if s is None: return ""
    if not isinstance(s, str): s = str(s)
    s = s.replace("\n", " ").replace("\r", " ").strip().lower()
    s = re.sub(r"\s+", " ", s)
    return s

def safe_int(x) -> int:
    if x is None: return 0
    if isinstance(x, (int, float)): return int(x)
    s = str(x).replace(" ", "").replace("\u00A0", "")
    s = re.sub(r"[^\d\-+]", "", s)
    try: return int(s) if s else 0
    except: return 0

def is_invoice_header(val) -> bool:
    if val is None: return False
    s = str(val).strip()
    if not s: return False
    s_norm = norm(s)
    if any(k in s_norm for k in ("надо отгруз", "надо отправ", "статус", "коммент", "сумма", "итог", "разниц")):
        return False
    if re.fullmatch(r"\d+", s):
        return True
    if re.fullmatch(r"\d{1,4}(?:[\-./][\wА-Яа-я]+)+", s):
        return True
    return False

# -------------------- size normalization --------------------

def normalize_size_token(s: str) -> str:
    if not s:
        return ""
    s0 = str(s).strip()
    s1 = s0.lower().replace("х", "x")
    s1 = re.sub(r"\s+", "", s1)
    s1 = re.sub(r"\s*-\s*", "-", s1)

    m = re.search(r"/([a-zx]+)$", s1)
    if m: s1 = m.group(1)

    m = re.fullmatch(r"([2-9])x*l", s1) or re.search(r"([2-9])x*l", s1)
    if m: return f"{m.group(1)}xl"

    m = re.search(r"^(x{2,})l$", s1)
    if m: return f"{len(m.group(1))}xl"

    if s1 in {"xxl"}:   return "2xl"
    if s1 in {"xxxl"}:  return "3xl"
    if s1 in {"xxxxl"}: return "4xl"
    if s1 in {"xs","s","m","l","xl"}: return s1

    if re.fullmatch(r"\d{2,3}(?:-\d{2,3}){1,3}", s1):
        return s1

    return s1

# -------------------- helpers --------------------

def find_cell_with_text(ws: Worksheet, text: str):
    target = norm(text)
    for r in ws.iter_rows(min_row=1, max_row=CONFIG["SCAN_MAX_ROWS"],
                          min_col=1, max_col=CONFIG["SCAN_MAX_COLS"]):
        for c in r:
            if target and target in norm(c.value):
                return c.row, c.column
    return None

def select_status_sheet(wb, debug: bool=False):
    """
    Выбор листа для статуса:
      1) видимый 'ПРОЕКТ'
      2) видимый 'ДЛЯ МОДЖИ'
      3) видимый лист с максимальным количеством совпадений по заголовкам:
         'артикул', 'размер', 'отгруз'/'отправ'
    """
    def visible(ws): return getattr(ws, "sheet_state", "visible") == "visible"

    # 1) "ПРОЕКТ"
    for ws in wb.worksheets:
        if visible(ws) and norm(ws.title) == norm("ПРОЕКТ"):
            if debug: log(f"  [DEBUG] Selected sheet by name: '{ws.title}' (ПРОЕКТ)")
            return ws
    # 2) "ДЛЯ МОДЖИ"
    for ws in wb.worksheets:
        if visible(ws) and norm(ws.title) == norm("ДЛЯ МОДЖИ"):
            if debug: log(f"  [DEBUG] Selected sheet by name: '{ws.title}' (ДЛЯ МОДЖИ)")
            return ws

    # 3) Автопоиск по заголовкам
    def sheet_score(ws: Worksheet) -> int:
        if not visible(ws): return -1
        score = 0
        max_r = min(CONFIG["SCAN_MAX_ROWS"], ws.max_row or 1)
        max_c = min(CONFIG["SCAN_MAX_COLS"], ws.max_column or 1)
        for r in range(1, max_r+1):
            for c in range(1, max_c+1):
                s = norm(ws.cell(row=r, column=c).value)
                if not s: continue
                if "артикул" in s: score += 2
                if "размер" in s:  score += 2
                if ("отгруз" in s) or ("отправ" in s): score += 2
                if "коммент" in s: score += 1
                if "сумма" in s or "итог" in s or "разниц" in s: score += 1
        return score

    scored = [(sheet_score(ws), ws) for ws in wb.worksheets]
    scored.sort(key=lambda x: x[0], reverse=True)
    if debug:
        log("  [DEBUG] Sheets and scores:")
        for sc, ws in scored:
            log(f"    - '{ws.title}' (state={getattr(ws,'sheet_state','visible')}): score={sc}")
    if scored and scored[0][0] >= 0:
        if debug: log(f"  [DEBUG] Selected sheet by score: '{scored[0][1].title}'")
        return scored[0][1]

    # fallback: первый лист
    ws0 = wb.worksheets[0] if wb.worksheets else None
    if debug and ws0: log(f"  [DEBUG] Selected sheet fallback: '{ws0.title}'")
    return ws0

def write_value_unmerged(ws: Worksheet, row:int, col:int, value):
    cell = ws.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        to_unmerge=None
        for rng in list(ws.merged_cells.ranges):
            if rng.min_row<=row<=rng.max_row and rng.min_col<=col<=rng.max_col:
                to_unmerge=rng; break
        if to_unmerge is not None:
            ws.unmerge_cells(range_string=str(to_unmerge))
    ws.cell(row=row, column=col).value = value

def copy_cell_style(dst_cell, src_cell):
    for attr in ("font", "fill", "border", "number_format", "alignment", "protection"):
        try:
            val = getattr(src_cell, attr, None)
            if val is not None:
                setattr(dst_cell, attr, copy(val))
        except Exception:
            pass

def clone_column_style(ws: Worksheet, src_col: int, dst_col: int, max_rows: Optional[int] = None):
    if max_rows is None: max_rows = ws.max_row
    try:
        ws.column_dimensions[get_column_letter(dst_col)].width = \
            ws.column_dimensions[get_column_letter(src_col)].width
    except Exception:
        pass
    for r in range(1, max_rows + 1):
        copy_cell_style(ws.cell(row=r, column=dst_col), ws.cell(row=r, column=src_col))

def estimate_last_data_row(ws: Worksheet,
                           header_row: int,
                           key_cols: List[int]) -> int:
    last = header_row
    bottom = ws.max_row
    for r in range(bottom, header_row, -1):
        for c in key_cols:
            v = ws.cell(row=r, column=c).value
            if v not in (None, ""):
                return r
    return last

# -------------------- read invoice --------------------

def read_invoice_number(wb) -> Optional[str]:
    candidates = ["Инвойс - Моджи", "Инвойс-Моджи", "Invoice - Moji", "Счет-фактура", "Счёт-фактура"]
    for name in candidates:
        for ws in wb.worksheets:
            if norm(ws.title) == norm(name):
                pos = (find_cell_with_text(ws,"№ инвойса")
                       or find_cell_with_text(ws,"№ инвойса:")
                       or find_cell_with_text(ws,"инвойс №"))
                if pos:
                    r,c = pos
                    val = ws.cell(row=r, column=c+1).value
                    return str(val).strip() if val is not None else None
    for ws in wb.worksheets:
        pos = (find_cell_with_text(ws,"№ инвойса")
               or find_cell_with_text(ws,"№ инвойса:")
               or find_cell_with_text(ws,"инвойс №"))
        if pos:
            r,c = pos
            val = ws.cell(row=r, column=c+1).value
            return str(val).strip() if val is not None else None
    return None

def find_header_row(ws: Worksheet, header_candidates: List[str]) -> Optional[int]:
    heads = [norm(h) for h in header_candidates]
    for r in ws.iter_rows(min_row=1, max_row=CONFIG["SCAN_MAX_ROWS"],
                          min_col=1, max_col=CONFIG["SCAN_MAX_COLS"]):
        texts = [norm(c.value) for c in r]
        if any(any(h in t for t in texts) for h in heads):
            return r[0].row
    return None

def find_col_in_row(ws: Worksheet, row_idx: int, candidates: List[str]) -> Optional[int]:
    heads = [norm(h) for h in candidates]
    for row in ws.iter_rows(min_row=row_idx, max_row=row_idx,
                            min_col=1, max_col=CONFIG["SCAN_MAX_COLS"]):
        for c in row:
            v = norm(c.value)
            if v in heads or any(h in v for h in heads):
                return c.column
    return None

def parse_items_from_invoice(wb)->Dict[Tuple[str,str],int]:
    ws = select_status_sheet(wb, debug=False) or wb.worksheets[0]
    if ws is None: return {}
    header_row = find_header_row(ws, ["Общее","Размер","Размеры","Цена"]) or 1
    col_article = find_col_in_row(ws, header_row, ["№ артикула","артикул модели","артикул"])
    col_price   = find_col_in_row(ws, header_row, ["цена","цена $","цена-"])
    col_total   = find_col_in_row(ws, header_row, ["общее","итого"])
    col_size_n  = find_col_in_row(ws, header_row, ["размер","размеры"])
    col_qty_n   = find_col_in_row(ws, header_row, ["заказ","кол-во","количество","qty","count"])

    items: Dict[Tuple[str,str],int] = {}

    if col_price and col_total:
        size_cols = []
        for c in range(col_price+1, col_total):
            header_val = ws.cell(row=header_row, column=c).value
            if header_val is None: continue
            size_cols.append((c, normalize_size_token(header_val)))
        last_row = ws.max_row
        for r in range(header_row+1, last_row+1):
            article = ws.cell(row=r, column=col_article).value if col_article else None
            article = str(article).strip() if article else ""
            if not article:
                continue
            for c,size_key in size_cols:
                q = safe_int(ws.cell(row=r, column=c).value)
                if q:
                    items[(article, size_key)] = items.get((article, size_key), 0) + q
        return items

    if col_size_n and col_qty_n:
        current_article = ""
        last_row = ws.max_row
        for r in range(header_row+1, last_row+1):
            a = ws.cell(row=r, column=col_article).value if col_article else None
            if a is not None and str(a).strip() != "":
                current_article = str(a).strip()
            size_name = ws.cell(row=r, column=col_size_n).value
            size_key = normalize_size_token(size_name)
            q = safe_int(ws.cell(row=r, column=col_qty_n).value)
            if not size_key or not q: continue
            items[(current_article, size_key)] = items.get((current_article, size_key), 0) + q
        return items

    return items

# -------------------- status update --------------------

def detect_headers(ws: Worksheet, debug=False):
    header_row = None
    col_article = col_size = col_ship = None
    for r in range(1, CONFIG["SCAN_MAX_ROWS"]+1):
        for c in range(1, CONFIG["SCAN_MAX_COLS"]+1):
            v = ws.cell(row=r, column=c).value
            if v is None: continue
            s = norm(v)
            if not s: continue
            if 'артикул' in s and col_article is None:
                col_article=c; header_row=r if header_row is None else min(header_row,r)
            if 'размер' in s and col_size is None:
                col_size=c; header_row=r if header_row is None else min(header_row,r)
            if ('отгруз' in s or 'отправ' in s) and col_ship is None:
                col_ship=c; header_row=r if header_row is None else min(header_row,r)
        if sum(x is not None for x in [col_article,col_size,col_ship])>=2:
            if header_row is None: header_row=r
            break
    if header_row is None: header_row=1
    if debug: log(f"  [DEBUG] header_row={header_row}, col_article={col_article}, col_size={col_size}, col_ship={col_ship}")
    return header_row, col_article, col_size, col_ship

def compute_invoice_area(ws: Worksheet, header_row: int, col_ship: int,
                         debug: bool=False) -> Tuple[int, int, Optional[int], Optional[int]]:
    max_col = ws.max_column
    def H(c): return norm(ws.cell(row=header_row, column=c).value)

    inv_cols_exist = [c for c in range(1, max_col+1) if is_invoice_header(ws.cell(row=header_row, column=c).value)]
    sum_cols = [c for c in range(col_ship+1, max_col+1) if any(tok in H(c) for tok in ("сумма","итог","разниц"))]
    right_anchors = [c for c in range(col_ship+1, max_col+1)
                     if any(tok in H(c) for tok in ("надо отгруз","надо отправ","статус","коммент","к погруз"))]

    if inv_cols_exist:
        inv_start = min(inv_cols_exist)
        stop_col  = max(inv_cols_exist) + 1
    else:
        if sum_cols:
            inv_start = max(sum_cols) + 1
            anchors_after = [c for c in right_anchors if c >= inv_start]
            stop_col = min(anchors_after) if anchors_after else (max_col + 1)
        else:
            inv_start = col_ship + 1
            anchors_after = [c for c in right_anchors if c >= inv_start]
            stop_col = min(anchors_after) if anchors_after else (max_col + 1)

    first_empty = None
    last_filled = None
    for c in range(inv_start, stop_col):
        v = ws.cell(row=header_row, column=c).value
        if v is None or str(v).strip() == "":
            if first_empty is None:
                first_empty = c
        else:
            last_filled = c

    if debug:
        log(f"  [DEBUG] invoice_area: inv_start={inv_start}, stop_col={stop_col}, "
            f"first_empty={first_empty}, last_filled={last_filled}, max_col={max_col}")

    return inv_start, stop_col, first_empty, last_filled

def insert_or_get_invoice_column(ws: Worksheet, header_row:int, invoice_no:str, col_ship:int, debug: bool=False)->int:
    invoice_no_s = str(invoice_no).strip()
    inv_start, stop_col, first_empty, last_filled = compute_invoice_area(ws, header_row, col_ship, debug=debug)

    for c in range(inv_start, stop_col):
        v = ws.cell(row=header_row, column=c).value
        if v is not None and str(v).strip().lower() == invoice_no_s.lower():
            return c

    template_col = last_filled if last_filled else inv_start
    last_data_row = estimate_last_data_row(ws, header_row, key_cols=[template_col, col_ship, template_col])

    if first_empty is not None:
        dst = first_empty
        # Распространяем стиль на весь лист, чтобы покрыть нижние итоги ("ИТОГО")
        clone_column_style(ws, template_col, dst, max_rows=ws.max_row)
        ws.cell(row=header_row, column=dst).value = invoice_no_s
        return dst

    insert_at = stop_col
    ws.insert_cols(insert_at, amount=1)
    # Распространяем стиль на весь лист, чтобы покрыть нижние итоги ("ИТОГО")
    clone_column_style(ws, template_col, insert_at, max_rows=ws.max_row)
    ws.cell(row=header_row, column=insert_at).value = invoice_no_s
    return insert_at

def update_shipment_formulas(ws: Worksheet, header_row:int, col_ship:int, debug: bool=False):
    inv_start, stop_col, first_empty, last_filled = compute_invoice_area(ws, header_row, col_ship, debug=debug)
    last = last_filled if last_filled is not None else inv_start
    if last < inv_start:
        return
    start_letter = get_column_letter(inv_start)
    end_letter   = get_column_letter(last)

    last_data_row = estimate_last_data_row(ws, header_row, key_cols=[col_ship, inv_start, last])

    for r in range(header_row+1, last_data_row+1):
        ws.cell(row=r, column=col_ship).value = f"=SUM({start_letter}{r}:{end_letter}{r})"

def update_invoice_totals(ws: Worksheet, header_row:int, col_size:int, invoice_col:int):
    col_letter = get_column_letter(invoice_col)
    sum_tokens = ("общее", "итог", "итого", "всего")

    def row_has_sum_token(r: int) -> bool:
        for c in range(1, col_size + 1):
            if any(t in norm(ws.cell(row=r, column=c).value) for t in sum_tokens):
                return True
        return False

    def row_is_size(r: int) -> bool:
        raw = ws.cell(row=r, column=col_size).value
        s = norm(raw)
        if not s:
            return False
        if any(t in s for t in sum_tokens):
            return False
        return True

    # Две границы: рабочая зона данных (для размера/колонок) и фактический низ листа
    last_data_row = estimate_last_data_row(ws, header_row, key_cols=[col_size, invoice_col])
    scan_last_row = ws.max_row

    group_start: Optional[int] = None
    last_size_row: Optional[int] = None
    per_total_rows: List[int] = []  # строки с промежуточными итогами по артикулам

    # Идём до реального низа листа, чтобы захватить последнюю строку "ОБЩЕЕ" и нижний "ИТОГО"
    for r in range(header_row + 1, scan_last_row + 1):
        if row_is_size(r):
            if group_start is None:
                group_start = r
            last_size_row = r
            continue

        if row_has_sum_token(r):
            if group_start is not None and last_size_row is not None and last_size_row >= group_start:
                ws.cell(row=r, column=invoice_col).value = f"=SUM({col_letter}{group_start}:{col_letter}{last_size_row})"
                per_total_rows.append(r)
            group_start = None
            last_size_row = None
            continue

    # Найти строку общего итога ("ИТОГО" предпочтительно, затем "ИТОГ", затем "ВСЕГО")
    def row_contains_token(r: int, token: str) -> bool:
        for c in range(1, col_size + 1):
            if token in norm(ws.cell(row=r, column=c).value):
                return True
        return False

    grand_row: Optional[int] = None
    for token in ("итого", "итог", "всего"):
        for r in range(scan_last_row, header_row, -1):
            if row_contains_token(r, token):
                grand_row = r
                break
        if grand_row is not None:
            break

    # Сформировать формулу суммирования всех промежуточных итогов в строке ИТОГО
    if grand_row is not None and per_total_rows:
        per_total_rows.sort()
        # сгруппировать подряд идущие строки в диапазоны для компактной формулы
        ranges = []
        start = end = per_total_rows[0]
        for r in per_total_rows[1:]:
            if r == end + 1:
                end = r
            else:
                ranges.append((start, end))
                start = end = r
        ranges.append((start, end))

        refs = []
        for a, b in ranges:
            if a == b:
                refs.append(f"{col_letter}{a}")
            else:
                refs.append(f"{col_letter}{a}:{col_letter}{b}")
        ws.cell(row=grand_row, column=invoice_col).value = f"=SUM({','.join(refs)})"

def fill_quantities(ws: Worksheet, header_row:int, invoice_col:int,
                    items:Dict[Tuple[str,str],int],
                    col_article:int, col_size:int)->int:
    index: Dict[Tuple[str,str],int] = {}
    current_article = ""

    last_data_row = estimate_last_data_row(ws, header_row, key_cols=[col_article, col_size])

    for r in range(header_row+1, last_data_row+1):
        a = ws.cell(row=r, column=col_article).value
        if a is not None and str(a).strip() != "":
            current_article = str(a).strip()
        s = ws.cell(row=r, column=col_size).value
        s_key = normalize_size_token(s)
        if current_article or s_key:
            index[(norm(current_article), s_key)] = r

    placed = 0
    for (article,size_key), qty in items.items():
        key = (norm(article), normalize_size_token(size_key))
        row = index.get(key)
        if row is None:
            continue
        ws.cell(row=row, column=invoice_col).value = safe_int(qty)
        placed += 1
    return placed

# -------------------- processed registry (persist) --------------------

def processed_registry_path(status_root: Path, factory_name: str) -> Path:
    return status_root / factory_name / CONFIG["OUTPUT_SUBFOLDER"] / "processed_invoices.txt"

def already_processed(invoice_path: Path, status_root: Path) -> bool:
    try:
        factory = invoice_path.parent.name
        reg = processed_registry_path(status_root, factory)
        if not reg.exists():
            return False
        mtime = int(Path(invoice_path).stat().st_mtime)
        key = f"{invoice_path.name}|{mtime}"
        with reg.open("r", encoding="utf-8") as f:
            for line in f:
                if line.strip() == key:
                    return True
        return False
    except Exception:
        return False

def mark_processed(invoice_path: Path, status_root: Path):
    try:
        factory = invoice_path.parent.name
        reg = processed_registry_path(status_root, factory)
        reg.parent.mkdir(exist_ok=True)
        mtime = int(Path(invoice_path).stat().st_mtime)
        key = f"{invoice_path.name}|{mtime}"
        with reg.open("a", encoding="utf-8") as f:
            f.write(key + "\n")
    except Exception:
        pass

# -------------------- xls -> xlsx conversion --------------------

def convert_xls_to_xlsx(xls_path: Path, debug: bool = False) -> Optional[Path]:
    """
    Convert .xls to .xlsx using Excel COM (pywin32). Returns path to .xlsx on success.
    On failure, returns None.
    """
    try:
        import win32com.client  # type: ignore
    except Exception:
        log("Для конвертации .xls установи pywin32:  pip install pywin32")
        return None

    xlsx_path = xls_path.with_suffix(".xlsx")
    try:
        if debug:
            log(f"  [DEBUG] Конвертация .xls -> .xlsx: {xls_path} -> {xlsx_path}")
        # удалить целевой, если уже есть
        if xlsx_path.exists():
            try:
                xlsx_path.unlink()
            except Exception as e:
                log(f"  [WARN] Не удалось удалить существующий {xlsx_path}: {e}")

        excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = False
        # дать файлу стабилизироваться
        time.sleep(max(0.5, CONFIG.get("FILE_SETTLE_SECONDS", 2)))
        wb = excel.Workbooks.Open(str(xls_path), UpdateLinks=0, ReadOnly=True)
        prev_alerts = excel.DisplayAlerts
        excel.DisplayAlerts = False
        try:
            wb.SaveAs(str(xlsx_path), FileFormat=51)  # 51 = xlOpenXMLWorkbook
        finally:
            excel.DisplayAlerts = prev_alerts
        wb.Close(False)
        try:
            excel.Quit()
        except Exception:
            pass

        # по желанию удалим исходник с задержкой
        if CONFIG.get("DELETE_XLS_AFTER_CONVERT", True):
            def _del(src=xls_path):
                try:
                    if src.exists():
                        src.unlink()
                        if debug:
                            log(f"  [DEBUG] Удалён исходный .xls: {src}")
                except Exception as e:
                    log(f"  [WARN] Не удалось удалить {src}: {e}")
            Timer(float(CONFIG.get("DELETE_XLS_DELAY", 5.0)), _del).start()

        return xlsx_path
    except Exception as e:
        log(f"  ! Ошибка конвертации {xls_path}: {e}")
        try:
            excel.Quit()  # type: ignore  # may not exist
        except Exception:
            pass
        return None

# -------------------- main processing --------------------

def process_invoice_file(invoice_path: Path, status_root: Path, dry_run=False, debug=False):
    log(f"[{datetime.now().strftime('%H:%M:%S')}] Обработка: {invoice_path}")
    factory_name = invoice_path.parent.name
    target_folder = status_root / factory_name
    if not target_folder.exists():
        log(f"  ! Не найден целевой каталог статусов: {target_folder}")
        return None

    out_dir = target_folder / CONFIG["OUTPUT_SUBFOLDER"]
    out_dir.mkdir(exist_ok=True)

    ready_candidates = [p for p in out_dir.glob("*.xls*")
                        if p.suffix.lower() in CONFIG["ALLOWED_STATUS_EXT"] and not p.name.startswith("~$")]
    if ready_candidates:
        status_file = max(ready_candidates, key=lambda p: p.stat().st_mtime)
    else:
        base_candidates = [p for p in target_folder.glob("*.xls*")
                           if p.suffix.lower() in CONFIG["ALLOWED_STATUS_EXT"] and not p.name.startswith("~$")]
        if not base_candidates:
            log(f"  ! В '{target_folder}' нет Excel-файлов для обновления.")
            return None
        status_file = max(base_candidates, key=lambda p: p.stat().st_mtime)

    if already_processed(invoice_path, status_root):
        log("  ↪ Пропуск: этот инвойс с тем же mtime уже обработан ранее.")
        return None

    # Если инвойс .xls — конвертируем во временный .xlsx рядом
    original_invoice_path = invoice_path
    if invoice_path.suffix.lower() == ".xls":
        converted = convert_xls_to_xlsx(invoice_path, debug=debug)
        if not converted:
            log("  ! Не удалось сконвертировать .xls в .xlsx (см. сообщение выше)")
            return None
        invoice_path = converted

    try:
        wb_inv = load_workbook(invoice_path, data_only=True, read_only=False)
    except Exception as e:
        log(f"  ! Не удалось открыть инвойс '{invoice_path.name}': {e}")
        return None

    inv_no = read_invoice_number(wb_inv)
    if not inv_no:
        log("  ! Не найден номер инвойса.")
        return None

    items = parse_items_from_invoice(wb_inv)
    if not items:
        log("  ! Не удалось извлечь позиции из инвойса (артикул/размер/кол-во).")
        return None

    log(f"  ➜ Инвойс № {inv_no}; позиций: {len(items)}")

    try:
        wb_stat = load_workbook(status_file, data_only=False, keep_vba=status_file.suffix.lower()=='.xlsm')
    except Exception as e:
        log(f"  ! Не удалось открыть файл статусов '{status_file.name}': {e}")
        return None

    ws = select_status_sheet(wb_stat, debug=debug)
    if ws is None:
        log("  ! Не найден подходящий лист со статусом.")
        return None
    if debug:
        log(f"  [DEBUG] Selected sheet: '{ws.title}'")

    header_row, col_article, col_size, col_ship = detect_headers(ws, debug=debug)
    if not col_ship or not col_size or not col_article:
        log("  ! Не удалось распознать ключевые колонки на листе статусов.")
        return None

    invoice_col = insert_or_get_invoice_column(ws, header_row, inv_no, col_ship, debug=debug)
    placed = fill_quantities(ws, header_row, invoice_col, items, col_article, col_size)
    update_invoice_totals(ws, header_row, col_size, invoice_col)
    update_shipment_formulas(ws, header_row, col_ship, debug=debug)
    if debug: log(f"  [DEBUG] Заполнено ячеек: {placed}; столбец инвойса: {invoice_col}")

    out_path = (target_folder / CONFIG["OUTPUT_SUBFOLDER"] / status_file.name)
    if dry_run:
        log(f"  [DRY RUN] Пропущено сохранение: {out_path}")
        mark_processed(original_invoice_path, status_root)
        return None

    delay = CONFIG["SAVE_RETRY_BASE_DELAY"]
    for attempt in range(1, CONFIG["SAVE_RETRIES"] + 1):
        try:
            wb_stat.save(out_path)
            log(f"  ✓ Сохранено: {out_path}")
            mark_processed(original_invoice_path, status_root)
            return out_path
        except PermissionError as e:
            log(f"  [WARN] Не удалось сохранить (попытка {attempt}/{CONFIG['SAVE_RETRIES']}): {e}")
            if attempt < CONFIG["SAVE_RETRIES"]:
                time.sleep(delay); delay *= 2; continue
            else:
                log(f"  ! Отказ сохранения после {CONFIG['SAVE_RETRIES']} попыток.")
                return None
        except Exception as e:
            log(f"  ! Не удалось сохранить файл: {e}")
            return None

def should_handle(p: Path)->bool:
    return p.suffix.lower() in CONFIG["ALLOWED_INVOICE_EXT"] and not p.name.startswith("~$")

# -------------------- watcher --------------------

class Handler(FileSystemEventHandler):
    def __init__(self, status_root: Path, dry_run: bool = False, debug: bool = False):
        super().__init__()
        self.status_root = status_root
        self.dry_run = dry_run
        self.debug = debug
        self.recent: Dict[str, Tuple[Tuple[int,int], float]] = {}

    def _file_signature(self, path: Path) -> Optional[Tuple[int,int]]:
        try:
            st = path.stat()
            return (int(st.st_size), int(st.st_mtime))
        except Exception:
            return None

    def _stable_signature(self, path: Path) -> Optional[Tuple[int,int]]:
        sig1 = self._file_signature(path)
        if sig1 is None:
            return None
        time.sleep(CONFIG["WATCH_STABLE_WAIT"])
        sig2 = self._file_signature(path)
        if sig1 == sig2:
            return sig1
        return None

    def _should_skip_by_recent(self, path: Path, sig: Tuple[int,int]) -> bool:
        now = time.time()
        cooldown = CONFIG["WATCH_DEDUP_COOLDOWN"]
        rec = self.recent.get(str(path))
        if rec:
            last_sig, ts = rec
            if last_sig == sig and (now - ts) < cooldown:
                if self.debug:
                    log(f"  [DEBUG] skip by recent cache (cooldown {cooldown}s): {path}")
                return True
        self.recent[str(path)] = (sig, now)
        return False

    def _handle_path(self, raw_path: str):
        p = Path(raw_path)
        if not should_handle(p):
            return
        time.sleep(CONFIG["FILE_SETTLE_SECONDS"])

        sig = self._stable_signature(p)
        if sig is None:
            time.sleep(CONFIG["FILE_SETTLE_SECONDS"])
            sig = self._file_signature(p)
            if sig is None:
                return

        # Если это .xls — сначала конвертируем и не обрабатываем здесь;
        # обработка произойдёт на событии создания .xlsx
        if p.suffix.lower() == ".xls":
            xlsx = convert_xls_to_xlsx(p, debug=self.debug)
            if self.debug:
                if xlsx:
                    log(f"  [DEBUG] Создан .xlsx: {xlsx}")
                else:
                    log("  [DEBUG] Конвертация .xls не удалась — обработка пропущена")
            return

        if self._should_skip_by_recent(p, sig):
            return

        if already_processed(p, self.status_root):
            if self.debug:
                log(f"  [DEBUG] already processed (persistent): {p}")
            return

        process_invoice_file(p, self.status_root, dry_run=self.dry_run, debug=self.debug)

    def on_created(self, event):
        if not event.is_directory:
            self._handle_path(event.src_path)

    def on_moved(self, event):
        if not event.is_directory:
            self._handle_path(event.dest_path)

# -------------------- CLI --------------------

def main():
    parser = argparse.ArgumentParser(description="Синхронизация инвойсов в статусы (умный выбор листа, адаптивная зона инвойсов).")
    parser.add_argument("--once", action="store_true", help="Единовременная обработка всех найденных инвойсов")
    parser.add_argument("--watch", action="store_true", help="Следить за новыми файлами в папках заводов")
    parser.add_argument("--invoices-root", default=CONFIG["INVOICES_ROOT"], help="Корень папок с инвойсами (по заводам)")
    parser.add_argument("--status-root",  default=CONFIG["STATUS_ROOT"],  help="Корень папок со статусами (по заводам)")
    parser.add_argument("--dry-run", action="store_true", help="Без сохранения результатов (только лог)")
    parser.add_argument("--debug", action="store_true", help="Расширенный лог")
    args = parser.parse_args()

    invoices_root = Path(args.invoices_root)
    status_root = Path(args.status_root)

    if args.once:
        log(f"Ищу инвойсы в: {invoices_root}")
        for factory_dir in invoices_root.iterdir():
            log(f"  Проверяю папку: {factory_dir}")
            if not factory_dir.is_dir():
                continue
            for f in sorted(factory_dir.glob("**/*")):
                if should_handle(f):
                    process_invoice_file(f, status_root, dry_run=args.dry_run, debug=args.debug)

    elif args.watch:
        if not WATCHDOG_AVAILABLE:
            log("Для режима --watch установи пакет watchdog:  pip install watchdog")
            return
        observer = Observer()
        handler = Handler(status_root=status_root, dry_run=args.dry_run, debug=args.debug)
        for factory_dir in invoices_root.iterdir():
            if factory_dir.is_dir():
                log(f"Наблюдение: {factory_dir}")
                observer.schedule(handler, str(factory_dir), recursive=False)
        observer.start()
        try:
            while True:
                time.sleep(1)
        except KeyboardInterrupt:
            observer.stop()
        observer.join()

    else:
        parser.print_help()

if __name__ == "__main__":
    main()
